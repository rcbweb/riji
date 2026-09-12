"""
report.py - figure + Excel workbook from the selected live cells.
=================================================================
Takes live_cells.csv (one row per SELECTED LIVE cell, written by rank_top) and
produces what a manuscript needs:

  uptake_figure.png / .pdf   one bar per experiment folder, every cell as a
                             point, mean +/- SD, significance brackets.
                             .pdf is vector, for figure assembly.
  uptake_box.png             box + points version (shows the distribution
                             shape, which a bar chart hides)
  Results.xlsx               Cells / For plotting / Summary / Statistics /
                             Per image / Per cell / Methods
  methods.txt                a methods paragraph with your actual numbers in it

The drawing itself lives in draw_bars(), which figure_editor.py also calls, so
the interactive preview and the exported file are literally the same code and
cannot drift apart.

Two levels of n
---------------
Cell-level  : n = cells. Sensitive, but cells from one image share a dish, a
              passage and an imaging session, so they are not independent.
Image-level : n = images, each contributing one mean. The conservative
              "biological replicate" analysis.
Both are always computed. level="image" plots the conservative one.

Run:  python -m cell_viability.report "<results folder>" [--metric=...]
      [--error=sd|sem|ci95] [--level=cell|image] [--title=...] [--no-stars]
"""
import os
import sys
import json
import itertools
import numpy as np

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from cell_viability import outputs as OUT  # noqa: E402

METRIC_LABEL = {
    "cy3_integrated": "Intensity (a.u.)",
    "cy3_mean": "Intensity per area (a.u.)",
}
DEFAULT_COLORS = ["#B3B3B3", "#7FA8C9", "#8FBF8F", "#D98C8C",
                  "#B39DDB", "#F0C05A", "#7FC9C4", "#E0A3C4"]

# A short, muted, print-safe palette. Offered as named choices so picking a
# colour is one click in a list rather than a modal RGB dialog.
PALETTE = [
    ("Grey", "#B3B3B3"), ("Dark grey", "#7A7A7A"), ("White", "#FFFFFF"),
    ("Blue", "#7FA8C9"), ("Green", "#8FBF8F"), ("Red", "#D98C8C"),
    ("Purple", "#B39DDB"), ("Orange", "#F0C05A"), ("Teal", "#7FC9C4"),
    ("Pink", "#E0A3C4"),
]
PALETTE_BY_NAME = dict(PALETTE)
PALETTE_BY_HEX = {v.upper(): k for k, v in PALETTE}


# ──────────────────────────────────────────────────────────────────────────
#  Data shaping
# ──────────────────────────────────────────────────────────────────────────

def _find_live_csv(path):
    path = os.path.abspath(path)
    if os.path.isfile(path):
        return path
    # data/ first: results written by the current layout. The older locations
    # are still accepted so an existing results folder keeps working, but they
    # must not win - a stale file from a previous layout silently produced a
    # figure from the wrong numbers.
    for cand in (os.path.join(path, "data", "live_cells.csv"),
                 os.path.join(path, "live_cells.csv"),
                 os.path.join(path, "_summary", "live_cells.csv"),
                 os.path.join(path, "viability_workspace", "top_cells",
                              "_summary", "live_cells.csv")):
        if os.path.isfile(cand):
            return cand
    raise SystemExit(
        f"Could not find live_cells.csv under {path}\n"
        "Point this at the results folder produced by the analysis.")


def _stars(p):
    return ("***" if p < 1e-3 else "**" if p < 1e-2 else
            "*" if p < 5e-2 else "ns")


def _wrap(label, width=14):
    words, lines, cur = str(label).split(), [], ""
    for w in words:
        if cur and len(cur) + 1 + len(w) > width:
            lines.append(cur)
            cur = w
        else:
            cur = f"{cur} {w}".strip()
    if cur:
        lines.append(cur)
    return "\n".join(lines)


def series_from_df(df, metric, level="cell", exclude_images=(),
                   normalize_to=None, order=None, include=None,
                   scope=OUT.SCOPE_TOP):
    """Turn the per-cell table into the arrays a bar chart needs.

    scope='top'        only the N cells per folder that were exported, so a
                       request for 20 produces 20 points
    scope='all'        every cell that passed the live-cell filter
    level='image'      each image contributes ONE mean (conservative n)
    exclude_images     {(condition, file)} or {file} to drop - for out-of-focus
                       fields, bubbles, wrong positions
    normalize_to       condition name; every value divided by that condition's
                       mean, so the axis becomes fold-change vs control
    Returns (conds, arrays, unit_note)

    The scope filter lives HERE, not in the callers. It used to be applied in
    build_report only, so the interactive editor drew every cell while the
    exported figure drew N - the same settings producing two different
    pictures. Anything that plots must go through this function.
    """
    d = df.copy()
    if (scope == OUT.SCOPE_TOP and "in_top_n" in d.columns
            and d["in_top_n"].sum()):
        d = d[d["in_top_n"] == 1]
    if exclude_images:
        ex_files = {e for e in exclude_images if not isinstance(e, tuple)}
        ex_pairs = {e for e in exclude_images if isinstance(e, tuple)}
        if ex_files:
            d = d[~d["file"].isin(ex_files)]
        if ex_pairs:
            key = list(zip(d["condition"], d["file"]))
            d = d[[k not in ex_pairs for k in key]]

    d = d[np.isfinite(d[metric])]
    conds = list(dict.fromkeys(d["condition"]))
    if include is not None:
        conds = [c for c in conds if c in include]
    if order:
        conds = [c for c in order if c in conds] + [c for c in conds if c not in order]

    arrays = []
    for c in conds:
        g = d[d["condition"] == c]
        v = (g.groupby("file")[metric].mean().to_numpy(float) if level == "image"
             else g[metric].to_numpy(float))
        arrays.append(v[np.isfinite(v)])

    note = ""
    if normalize_to and normalize_to in conds:
        base = arrays[conds.index(normalize_to)]
        b = float(base.mean()) if len(base) else 0.0
        if b > 0:
            arrays = [a / b for a in arrays]
            note = f"fold vs {normalize_to}"
    return conds, arrays, note


def _summarise(df, metric):
    rows = []
    for cond, g in df.groupby("condition", sort=False):
        v = g[metric].to_numpy(float)
        v = v[np.isfinite(v)]
        im = g.groupby("file")[metric].mean().to_numpy(float)
        im = im[np.isfinite(im)]
        n = len(v)
        sd = float(v.std(ddof=1)) if n > 1 else 0.0
        sem = sd / np.sqrt(n) if n else 0.0
        rows.append(dict(
            condition=cond, n_cells=n, n_images=int(g["file"].nunique()),
            mean=float(v.mean()) if n else 0.0, sd=sd, sem=sem,
            ci95_lo=float(v.mean() - 1.96 * sem) if n else 0.0,
            ci95_hi=float(v.mean() + 1.96 * sem) if n else 0.0,
            median=float(np.median(v)) if n else 0.0,
            q25=float(np.percentile(v, 25)) if n else 0.0,
            q75=float(np.percentile(v, 75)) if n else 0.0,
            min=float(v.min()) if n else 0.0, max=float(v.max()) if n else 0.0,
            mean_of_image_means=float(im.mean()) if len(im) else 0.0,
            sd_of_image_means=float(im.std(ddof=1)) if len(im) > 1 else 0.0,
        ))
    return rows


def _pairwise(df, metric):
    from scipy.stats import mannwhitneyu
    conds = list(dict.fromkeys(df["condition"]))
    cells = {c: df.loc[df.condition == c, metric].to_numpy(float) for c in conds}
    cells = {c: v[np.isfinite(v)] for c, v in cells.items()}
    imgs = {c: df[df.condition == c].groupby("file")[metric].mean().to_numpy(float)
            for c in conds}
    pairs = list(itertools.combinations(conds, 2))
    n_comp = max(1, len(pairs))
    out = []
    for a, b in pairs:
        ca, cb = cells[a], cells[b]
        if len(ca) < 3 or len(cb) < 3:
            continue
        u, p = mannwhitneyu(ca, cb, alternative="two-sided")
        pb = min(1.0, p * n_comp)
        ia, ib = imgs[a], imgs[b]
        if len(ia) >= 3 and len(ib) >= 3:
            ui, pi = mannwhitneyu(ia, ib, alternative="two-sided")
            pib = min(1.0, pi * n_comp)
        else:
            ui = pi = pib = float("nan")
        out.append(dict(
            condition_a=a, condition_b=b,
            n_cells_a=len(ca), n_cells_b=len(cb),
            median_a=float(np.median(ca)), median_b=float(np.median(cb)),
            fold_b_over_a=(float(np.median(cb) / np.median(ca))
                           if np.median(ca) else float("nan")),
            U_cells=float(u), p_cells=float(p), p_cells_bonferroni=float(pb),
            stars_cells=_stars(pb),
            n_images_a=len(ia), n_images_b=len(ib),
            U_images=float(ui), p_images=float(pi), p_images_bonferroni=float(pib),
            stars_images=_stars(pib) if pib == pib else "n/a",
        ))
    return out, n_comp


# ──────────────────────────────────────────────────────────────────────────
#  Drawing - shared by the exporter AND the interactive editor
# ──────────────────────────────────────────────────────────────────────────

def draw_bars(ax, conds, data, labels=None, colors=None, error="sd",
              show_points=True, show_stars=True, show_n=False,
              title="", xlabel="", ylabel=None, metric="cy3_integrated",
              unit_note="", ymax=None, point_size=13, tick_max=None):
    """Render the publication bar chart onto `ax`. Returns the y-scale exponent.

    tick_max   the largest value the measurement can actually take. The axis
               still stretches above it to hold the significance brackets, but
               no tick is labelled past it. Uptake has no such bound and
               passes None, which is exactly the behaviour this always had.
    """
    from scipy.stats import mannwhitneyu

    keep = [i for i, v in enumerate(data) if len(v) >= 1]
    conds = [conds[i] for i in keep]
    data = [data[i] for i in keep]
    if labels:
        labels = [labels[i] for i in keep]
    if not data:
        ax.text(0.5, 0.5, "No data to plot", ha="center", va="center")
        ax.axis("off")
        return 0

    # Fold magnitude into the axis label rather than leaving a floating "1e7".
    vmax = max(float(v.max()) for v in data) if data else 0.0
    exp = 0
    if vmax > 0 and not unit_note:
        e = int(np.floor(np.log10(vmax)))
        exp = 3 * (e // 3)
    scale = 10.0 ** exp
    if exp:
        data = [v / scale for v in data]

    means = np.array([v.mean() for v in data])
    sds = np.array([v.std(ddof=1) if len(v) > 1 else 0.0 for v in data])
    sems = np.array([s / np.sqrt(max(1, len(v))) for s, v in zip(sds, data)])
    err = {"sd": sds, "sem": sems, "ci95": 1.96 * sems}.get(error, sds)

    cols = colors or [DEFAULT_COLORS[i % len(DEFAULT_COLORS)] for i in range(len(data))]
    x = np.arange(len(data))

    ax.bar(x, means, width=0.62, color=cols, edgecolor="black",
           linewidth=1.6, zorder=2)
    ax.errorbar(x, means, yerr=err, fmt="none", ecolor="black",
                elinewidth=1.6, capsize=6, capthick=1.6, zorder=3)

    if show_points:
        rng = np.random.default_rng(0)
        for i, v in enumerate(data):
            ax.scatter(x[i] + rng.uniform(-0.17, 0.17, len(v)), v,
                       s=point_size, color="black", zorder=4, linewidths=0)

    disp = labels if labels else conds
    ax.set_xticks(x)
    ax.set_xticklabels([_wrap(c) for c in disp], fontsize=11, fontweight="bold")

    if ylabel is None:
        base = METRIC_LABEL.get(metric, "Intensity (a.u.)")
        if unit_note:
            ylabel = f"Uptake ({unit_note})"
        elif exp:
            ylabel = base.replace("(a.u.)", r"(a.u. $\times 10^{%d}$)" % exp)
        else:
            ylabel = base
    ax.set_ylabel(ylabel, fontsize=15, fontweight="bold", labelpad=8)
    if xlabel:
        ax.set_xlabel(xlabel, fontsize=15, fontweight="bold", labelpad=8)
    if title:
        ax.set_title(title, fontsize=16, fontweight="bold", pad=12)

    ax.tick_params(axis="both", which="major", direction="out",
                   length=6, width=1.6, labelsize=12)
    for lbl in ax.get_yticklabels():
        lbl.set_fontweight("bold")
    for sp in ax.spines.values():
        sp.set_linewidth(1.6)
    ax.set_ylim(bottom=0)
    ax.margins(x=0.08)

    ceiling = max(vmax / scale if scale else vmax, float((means + err).max()))

    if show_stars and len(data) > 1:
        pairs = list(itertools.combinations(range(len(data)), 2))
        n_comp = max(1, len(pairs))
        sig = []
        for a, b in pairs:
            if len(data[a]) < 3 or len(data[b]) < 3:
                continue
            _u, p = mannwhitneyu(data[a], data[b], alternative="two-sided")
            sig.append((a, b, min(1.0, p * n_comp)))
        step = ceiling * 0.085
        lvl = ceiling * 1.06
        for k, (a, b, p) in enumerate(sorted(sig, key=lambda s: abs(s[0] - s[1]))):
            yy = lvl + k * step * 1.5
            ax.plot([a, a, b, b], [yy - step * .28, yy, yy, yy - step * .28],
                    lw=1.4, color="black", clip_on=False)
            ax.text((a + b) / 2, yy + step * .06, _stars(p), ha="center",
                    va="bottom", fontsize=13, fontweight="bold")
        if sig:
            ceiling = lvl + (len(sig) - 1) * step * 1.5 + step

    if show_n:
        for i, v in enumerate(data):
            ax.text(x[i], 0, f"n={len(v)}", ha="center", va="bottom",
                    fontsize=9, color="#222")

    # The floor follows the data instead of being nailed to zero. Uptake is
    # never negative so this is still 0 for the cell figures, but a Pearson r
    # runs -1..+1 and a hard zero floor quietly cropped every anti-correlated
    # cell off the bottom of the chart - 27 of 171 in one real dataset, gone
    # from the picture while still counted in the mean.
    # Keyed on the measured values only, NOT on mean-minus-error. An error bar
    # reaching below zero on a quantity that cannot be negative is an artefact
    # of drawing it symmetrically, and letting that drag the axis down would
    # change every uptake figure. Real negative data is a different matter.
    floor = min(0.0, min((float(v.min()) for v in data if len(v)), default=0.0))
    ax.set_ylim(floor, float(ymax) if ymax else ceiling * 1.10)

    if tick_max is not None:
        # A bounded measurement. The axis has to reach past the maximum to
        # hold the brackets, but labelling ticks up there advertises values
        # the measurement cannot take: M1 is a fraction, there is no 1.25 of
        # it. Keep the room, drop the labels.
        keep_t = [t for t in ax.get_yticks()
                  if floor - 1e-9 <= t <= float(tick_max) + 1e-9]
        if keep_t:
            ax.set_yticks(keep_t)
            ax.set_ylim(floor, float(ymax) if ymax else ceiling * 1.10)
    return exp


def _rcparams():
    return {"font.family": "sans-serif",
            "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
            "axes.linewidth": 1.6}


def make_figure(conds, data, out_base, unit_note="", **kw):
    """Export the bar figure to PNG (300 dpi) + PDF (vector)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update(_rcparams())
    fig, ax = plt.subplots(figsize=(0.95 * max(1, len(conds)) + 2.0, 4.9))
    draw_bars(ax, conds, data, unit_note=unit_note, **kw)
    fig.tight_layout()
    fig.savefig(out_base + ".png", dpi=300, bbox_inches="tight")
    fig.savefig(out_base + ".pdf", bbox_inches="tight")
    plt.close(fig)
    print(f"[report] figure -> {out_base}.png")
    print(f"[report] figure -> {out_base}.pdf  (vector)")
    return out_base + ".png"


def make_boxplot(conds, data, out_png, metric="cy3_integrated", title=""):
    """Box + points. A bar chart hides distribution shape; this shows it."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update(_rcparams())
    data = [v for v in data if len(v)]
    if len(data) < 1:
        return None
    fig, ax = plt.subplots(figsize=(0.95 * len(data) + 2.0, 4.9))
    pos = np.arange(1, len(data) + 1)
    bp = ax.boxplot(data, positions=pos, widths=0.55, showfliers=False,
                    patch_artist=True, medianprops=dict(color="black", lw=1.6))
    for i, b in enumerate(bp["boxes"]):
        b.set(facecolor=DEFAULT_COLORS[i % len(DEFAULT_COLORS)],
              alpha=0.45, edgecolor="black", linewidth=1.4)
    rng = np.random.default_rng(0)
    for i, v in enumerate(data):
        ax.scatter(pos[i] + rng.uniform(-0.16, 0.16, len(v)), v, s=11,
                   color="black", alpha=0.65, zorder=3, linewidths=0)
    ax.set_yscale("log")
    ax.set_ylabel(METRIC_LABEL.get(metric, "Intensity (a.u.)"),
                  fontsize=13, fontweight="bold")
    ax.set_xticks(pos)
    ax.set_xticklabels([_wrap(c) for c in conds[:len(data)]], fontsize=10)
    if title:
        ax.set_title(title, fontsize=14, fontweight="bold")
    ax.grid(axis="y", ls=":", color="#ccc", alpha=0.7)
    fig.tight_layout()
    fig.savefig(out_png, dpi=200, bbox_inches="tight")
    plt.close(fig)
    return out_png


# ──────────────────────────────────────────────────────────────────────────
#  Methods paragraph
# ──────────────────────────────────────────────────────────────────────────

def methods_text(summary, stats, n_comp, metric, error, level, picker_info=None):
    conds = [r for r in summary if r["n_cells"]]
    per = "; ".join(f"{r['condition']}: n = {r['n_cells']} cells from "
                    f"{r['n_images']} images" for r in conds)
    pick = ""
    if picker_info:
        pick = (f" The selector was trained on {picker_info.get('n_positive','?')} "
                f"manually annotated cells and validated by image-grouped "
                f"cross-validation (AUC = {picker_info.get('cv_auc', float('nan')):.3f}), "
                f"so no image contributed to both training and testing.")
    errname = {"sd": "standard deviation", "sem": "standard error of the mean",
               "ci95": "95% confidence interval"}.get(error, "standard deviation")
    unit = "integrated over the whole cell" if metric == "cy3_integrated" \
        else "averaged over the cell area"
    txt = f"""METHODS (auto-generated - check the wording, the numbers are yours)

Image analysis. Cells were segmented from the transmitted-light (brightfield)
channel using Cellpose-SAM; the fluorescence channel was not used for
segmentation, so cell identification was independent of the signal being
quantified. Segmented objects were filtered to viable, well-separated,
completely imaged cells using a logistic-regression classifier built on
brightfield morphology, refractility, texture and neighbour-isolation
features.{pick} Objects that were rounded and phase-bright (characteristic of
dying cells), in contact with neighbours, or clipped by the field of view were
excluded before any fluorescence measurement.

Quantification. For each retained cell, fluorescence intensity was {unit}
after subtracting the median intensity of a local background annulus
surrounding that cell. Values are reported in arbitrary units and are
comparable between conditions acquired with identical settings.

Statistics. Data are presented as mean +/- {errname}, with each point
representing one {'image' if level == 'image' else 'cell'} ({per}). Because
uptake distributions are right-skewed, comparisons used two-sided Mann-Whitney
U tests with Bonferroni correction for {n_comp} pairwise comparison(s).
Significance is denoted *** p < 0.001, ** p < 0.01, * p < 0.05, ns = not
significant.

RESULTS SENTENCES (fill in as appropriate)
"""
    for s in stats:
        txt += (f"  {s['condition_a']} vs {s['condition_b']}: "
                f"{s['fold_b_over_a']:.2f}-fold, "
                f"cell-level p = {s['p_cells_bonferroni']:.2e} ({s['stars_cells']}), "
                f"image-level p = {s['p_images_bonferroni']:.2e} "
                f"({s['stars_images']}).\n")
    txt += """
NOTE ON n. Cell-level tests treat every cell as independent. Cells imaged in
the same field share a dish, passage and session, so the image-level p-value
is the conservative one. Where the two disagree, report the image-level result
or add imaging replicates.
"""
    return txt


# ──────────────────────────────────────────────────────────────────────────
#  Excel
# ──────────────────────────────────────────────────────────────────────────

def _tables_from_cells(df, metric):
    """Rebuild the per-folder "# | Cell | Intensity" listing from the measured
    cells, so the workbook is complete no matter which stage writes it."""
    import numpy as _np
    out = {}
    d = df[_np.isfinite(df[metric]) & (df[metric] > 0)]
    if "in_top_n" in d.columns and d["in_top_n"].sum():
        d = d[d["in_top_n"] == 1]
    for cond, g in d.groupby("condition"):
        g = g.sort_values(metric, ascending=False)
        rows = []
        for i, (_, r) in enumerate(g.iterrows(), 1):
            stem = os.path.splitext(str(r.get("file", "")))[0]
            name = f"{stem}  x{int(r.get('x', 0))} y{int(r.get('y', 0))}"
            rows.append((i, name, float(r[metric]),
                         int(r.get("passed_filter", 1))))
        if rows:
            out[cond] = rows
    return out


def write_excel(df, metric, summary, stats, n_comp, out_xlsx, error="sd",
                methods="", simple_tables=None, summary_full=None,
                scope="top", stats_full_df=None):
    import pandas as pd
    from openpyxl.styles import (Font, Alignment, PatternFill, Border,
                                 Side)
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.error_bar import ErrorBars
    from openpyxl.chart.data_source import NumDataSource, NumRef

    HDR_FILL = PatternFill("solid", fgColor="2C3E50")
    BAND_FILL = PatternFill("solid", fgColor="F4F7FA")
    _thin = Side(style="thin", color="D0D7DE")
    BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

    sum_df = pd.DataFrame(summary)
    stats_df = pd.DataFrame(stats)
    per_img = (df.groupby(["condition", "file"])[metric]
                 .agg(n_cells="size", mean="mean", median="median",
                      sd=lambda s: s.std(ddof=1)).reset_index())
    cell_cols = [c for c in ["condition", "file", "x", "y", "area_um2",
                             "cy3_integrated", "cy3_mean", "cy3_raw_mean",
                             "cy3_raw_integrated", "cy3_background",
                             "n_pixels", "pick_score", "uptake_rank"]
                 if c in df.columns]
    per_cell = df[cell_cols].copy()
    
    # Rename columns to match Fiji conventions
    # One naming table for every sheet. Units are part of the header, because
    # a column called "mean" tells the reader nothing about what it measures.
    rename_map = {
        # measurements
        "cy3_integrated": "Corrected IntDen (AU)",
        "cy3_mean": "Corrected Mean (AU)",
        "cy3_raw_mean": "Mean (AU)",
        "cy3_raw_integrated": "IntDen (AU)",
        "cy3_background": "Background (AU)",
        "area_um2": "Area (um2)",
        "n_pixels": "Area (pixels)",
        "mean": "Mean (AU)",
        "median": "Median (AU)",
        "sd": "SD (AU)",
        "sem": "SEM (AU)",
        "ci95_lo": "95% CI low (AU)",
        "ci95_hi": "95% CI high (AU)",
        "q25": "25th percentile (AU)",
        "q75": "75th percentile (AU)",
        "min": "Min (AU)",
        "max": "Max (AU)",
        "mean_of_image_means": "Mean of image means (AU)",
        "sd_of_image_means": "SD of image means (AU)",
        # identity and provenance
        "condition": "Folder",
        "file": "Image",
        "x": "x (pixels)",
        "y": "y (pixels)",
        "n_cells": "Cells (n)",
        "n_images": "Images (n)",
        "pick_score": "Confidence (0-1)",
        "uptake_rank": "Rank in folder",
        "passed_filter": "Passed filter",
        "decided_by": "Chosen by",
        # statistics
        "condition_a": "Folder A",
        "condition_b": "Folder B",
        "n_cells_a": "Cells in A (n)",
        "n_cells_b": "Cells in B (n)",
        "n_images_a": "Images in A (n)",
        "n_images_b": "Images in B (n)",
        "median_a": "Median A (AU)",
        "median_b": "Median B (AU)",
        "fold_b_over_a": "Fold (B / A)",
        "U_cells": "Mann-Whitney U (per cell)",
        "p_cells": "p (per cell)",
        "p_cells_bonferroni": "p (per cell, corrected)",
        "stars_cells": "Significance (per cell)",
        "U_images": "Mann-Whitney U (per image)",
        "p_images": "p (per image)",
        "p_images_bonferroni": "p (per image, corrected)",
        "stars_images": "Significance (per image)",
    }

    sum_df = sum_df.rename(columns=rename_map)
    stats_df = stats_df.rename(columns=rename_map)
    per_img = per_img.rename(columns=rename_map)
    per_cell = per_cell.rename(columns=rename_map)
    readme = pd.DataFrame({"": (methods or "").split("\n")})

    # One workbook, written in one pass. The intensities used to be a second
    # file appended here, which left duplicate sheets saying the same thing;
    # they are now the "Cells" and "For plotting" sheets below.
    kwargs = {"mode": "w"}

    # "Cells" is the sheet people actually read. It used to be built only when
    # rank_top passed the table in, so re-exporting from the figure editor
    # silently dropped it - editing a title removed the main sheet from the
    # workbook. Rebuild it from the data when it is not supplied.
    if not simple_tables:
        simple_tables = _tables_from_cells(df, metric)

    # The header carries the unit, so a column of large numbers is not left
    # for the reader to guess at.
    unit_col = ("Corrected IntDen (AU)" if metric == "cy3_integrated"
                else "Corrected Mean (AU)")
    cells_rows = []
    for cond, rows in (simple_tables or {}).items():
        for item in rows:
            rk, nm, val, ok = (item if len(item) == 4
                               else (item[0], "", item[1], item[2]))
            cells_rows.append({"#": rk, "Folder": cond, "Cell": nm,
                               unit_col: round(float(val), 1),
                               "Passed filter": "yes" if ok else "no (topped up)"})
    cells_df = pd.DataFrame(cells_rows)

    wide = {c: [round(float(r[2] if len(r) == 4 else r[1]), 1) for r in rows]
            for c, rows in (simple_tables or {}).items()}
    longest = max((len(v) for v in wide.values()), default=0)
    # side-by-side sheet: folders as columns, so it can be selected and charted
    wide_df = pd.DataFrame(dict({"#": list(range(1, longest + 1))},
                                **{c: v + [None] * (longest - len(v))
                                   for c, v in wide.items()}))

    with pd.ExcelWriter(out_xlsx, engine="openpyxl", **kwargs) as xl:
        if len(cells_df):
            cells_df.to_excel(xl, sheet_name="Cells", index=False)
        if longest:
            wide_df.to_excel(xl, sheet_name="For plotting", index=False,
                             startrow=1)
        sum_df.to_excel(xl, sheet_name="Summary", index=False, startrow=1)
        stats_df.to_excel(xl, sheet_name="Statistics", index=False)
        per_img.to_excel(xl, sheet_name="Per image", index=False)
        per_cell.to_excel(xl, sheet_name="Per cell", index=False)
        # these two were written straight from the raw frames, so they were the
        # only sheets still showing internal names like n_cells_a
        if summary_full is not None:
            pd.DataFrame(summary_full).rename(columns=rename_map).to_excel(
                xl, sheet_name="All selected cells", index=False)
        if stats_full_df is not None and len(stats_full_df):
            stats_full_df.rename(columns=rename_map).to_excel(
                xl, sheet_name="Statistics (all cells)", index=False)
        readme.to_excel(xl, sheet_name="Methods", index=False, header=False)

        ws = xl.sheets["Summary"]
        ws["A1"] = f"Live-cell uptake summary - {METRIC_LABEL.get(metric, metric)}"
        ws["A1"].font = Font(bold=True, size=13)
        if "For plotting" in xl.sheets:
            wp = xl.sheets["For plotting"]
            wp["A1"] = ("One column per folder - "
                        f"{METRIC_LABEL.get(metric, metric)}. "
                        "Select a block and insert a chart.")
            wp["A1"].font = Font(bold=True, size=11)

        def _number_format(header):
            """Pick a format from what the column measures."""
            h = str(header)
            if h.startswith("p (") or h.startswith("p_"):
                return "0.00E+00"          # probabilities span many decades
            if "(AU)" in h or "(pixels)" in h:
                return "#,##0"             # intensities are large; no decimals
            if "(um2)" in h:
                return "#,##0.0"
            if h.startswith("Fold") or h.startswith("Confidence"):
                return "0.00"
            if "(n)" in h or h in ("#", "Rank in folder"):
                return "0"
            if h.startswith("Mann-Whitney"):
                return "#,##0"
            return None

        def _style(sh, hdr_row):
            """Dark header, banded rows, borders, units-aware numbers."""
            head = [c.value for c in sh[hdr_row]]
            ncol = len([h for h in head if h is not None])
            for c in sh[hdr_row]:
                if c.value is None:
                    continue
                c.font = Font(bold=True, color="FFFFFF", size=10)
                c.fill = HDR_FILL
                c.alignment = Alignment(horizontal="center", vertical="center",
                                        wrap_text=True)
                c.border = BOX
            sh.row_dimensions[hdr_row].height = 30

            fmts = [_number_format(h) for h in head]
            for i, row in enumerate(sh.iter_rows(min_row=hdr_row + 1,
                                                 max_col=max(1, ncol))):
                for c in row:
                    c.border = BOX
                    if i % 2:
                        c.fill = BAND_FILL
                    j = c.column - 1
                    if j < len(fmts) and fmts[j] and isinstance(c.value, (int, float)):
                        c.number_format = fmts[j]
            for col in sh.columns:
                letter = get_column_letter(col[0].column)
                w = max((len(str(c.value)) for c in col[:60] if c.value is not None),
                        default=10)
                sh.column_dimensions[letter].width = min(30, max(11, w + 3))
            sh.freeze_panes = sh.cell(row=hdr_row + 1, column=1)
            if sh.max_row > hdr_row + 3 and ncol:
                sh.auto_filter.ref = (f"A{hdr_row}:"
                                      f"{get_column_letter(ncol)}{sh.max_row}")

        for name in ("Cells", "For plotting", "Summary", "Statistics",
                     "Statistics (all cells)", "Per image", "Per cell",
                     "All selected cells"):
            if name in xl.sheets:
                _style(xl.sheets[name],
                       2 if name in ("Summary", "For plotting") else 1)

        xl.sheets["Methods"].column_dimensions["A"].width = 90

        n = len(sum_df)
        hdr, first, last = 2, 3, 2 + n
        cols = {c: i + 1 for i, c in enumerate(sum_df.columns)}
        ch = BarChart()
        ch.type, ch.legend = "col", None
        ch.title = "Live-cell uptake"
        ch.y_axis.title = METRIC_LABEL.get(metric, "Intensity (a.u.)")
        ch.x_axis.title = "Condition"
        ch.height, ch.width = 9, 16
        # Look the columns up by their renamed headers. Hard-coding the old
        # lowercase names meant renaming a column silently broke this chart.
        ch.add_data(Reference(ws, min_col=cols[rename_map["mean"]],
                              min_row=hdr, max_row=last),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=cols[rename_map["condition"]],
                                    min_row=first, max_row=last))
        ec = get_column_letter(cols[rename_map["sd" if error == "sd" else "sem"]])
        ref = NumRef(f=f"'Summary'!${ec}${first}:${ec}${last}")
        ch.series[0].errBars = ErrorBars(
            errDir="y", errBarType="both", errValType="cust",
            plus=NumDataSource(numRef=ref), minus=NumDataSource(numRef=ref))
        ws.add_chart(ch, f"A{last + 3}")

    print(f"[report] workbook -> {out_xlsx}")
    return out_xlsx


# ──────────────────────────────────────────────────────────────────────────
#  Entry point
# ──────────────────────────────────────────────────────────────────────────

def build_report(path, outdir=None, metric="cy3_integrated", error="sd",
                 level="cell", title="", xlabel="", show_stars=True,
                 settings=None, simple_tables=None, plot_scope="top"):
    import pandas as pd

    live_csv = OUT.find_live_cells(path)
    # Figures and tables belong in ONE place, so whatever a caller passes -
    # the results folder, the figures folder, or nothing at all - is resolved
    # to the same folder here. The editor passed the results root and this
    # wrote straight into it, so an edited figure landed loose beside the
    # folder it belonged in while the dialog named the folder. Deciding it
    # here means no caller can put the deliverables somewhere else.
    outdir = OUT.figures_dir(OUT.results_root(outdir or live_csv), create=True)
    root_out = OUT.results_root(outdir)
    OUT.data_dir(root_out, create=True)
    # A results folder written by an older build has a stale figure and
    # workbook loose at the top level. Move them aside before writing the new
    # ones, so the folder holds one answer rather than two.
    OUT.retire_stale(root_out, verbose=False)

    df = pd.read_csv(live_csv)
    if metric not in df.columns:
        raise SystemExit(f"{metric} not in {live_csv}")

    st = settings or {}
    # Scope is applied inside series_from_df, so the exported figure and the
    # interactive preview cannot diverge. `full` keeps the whole population for
    # the comparison statistics.
    scope = st.get("plot_scope", plot_scope)
    full = df.copy()
    has_signal = (df.groupby("condition")[metric]
                    .apply(lambda s: np.isfinite(s).any() and (s.fillna(0) > 0).any()))
    plot_df = df[df["condition"].isin(has_signal[has_signal].index)].copy()
    dropped = [c for c in has_signal.index if not has_signal[c]]

    order = st.get("order") or (plot_df.groupby("condition")[metric].median()
                                .sort_values(ascending=False).index.tolist())
    conds, data, note = series_from_df(
        plot_df, metric, scope=scope, level=st.get("level", level),
        exclude_images=set(map(tuple, st["exclude"])) if st.get("exclude") else (),
        normalize_to=st.get("normalize_to"), order=order,
        include=st.get("include"))

    plot_df["condition"] = pd.Categorical(plot_df["condition"],
                                          [c for c in order
                                           if c in set(plot_df["condition"])],
                                          ordered=True)
    plot_df = plot_df.sort_values("condition")

    shown = df
    if (scope == OUT.SCOPE_TOP and "in_top_n" in df.columns
            and df["in_top_n"].sum()):
        shown = df[df["in_top_n"] == 1]
    summary = _summarise(shown, metric)
    summary_full = _summarise(full, metric)
    # The figure shows the top N per folder, but a p-value from the brightest
    # tail of each group is not the same claim as one from the whole
    # population, and the two can disagree. Both are computed, and any
    # disagreement is reported loudly rather than left to be discovered.
    stats_full, _nc_full = _pairwise(full, metric)
    stats_src = plot_df
    if (scope == OUT.SCOPE_TOP and "in_top_n" in plot_df.columns
            and plot_df["in_top_n"].sum()):
        stats_src = plot_df[plot_df["in_top_n"] == 1]
    stats, n_comp = _pairwise(stats_src, metric)

    picker = None
    for cand in (os.path.join(os.path.dirname(live_csv), "..", "..",
                              "pick_model.json"),
                 os.path.join(os.path.dirname(os.path.abspath(__file__)),
                              "default_pick_model.json")):
        if os.path.isfile(cand):
            try:
                with open(cand, encoding="utf-8") as fh:
                    picker = json.load(fh)
                break
            except Exception:
                pass

    # Summary and Statistics are sheets in Results.xlsx. Writing them again as
    # CSVs only created second copies that could fall out of step with the
    # workbook after an edit.

    meth = methods_text(summary, stats, n_comp, metric, st.get("error", error),
                        st.get("level", level), picker)
    with open(os.path.join(outdir, "methods.txt"), "w", encoding="utf-8") as fh:
        fh.write(meth)

    fig = make_figure(conds, data, os.path.join(outdir, "uptake_figure"),
                      labels=[st.get("labels", {}).get(c, c) for c in conds],
                      colors=[st.get("colors", {}).get(c) or
                              DEFAULT_COLORS[i % len(DEFAULT_COLORS)]
                              for i, c in enumerate(conds)],
                      error=st.get("error", error),
                      show_points=st.get("show_points", True),
                      show_stars=st.get("show_stars", show_stars),
                      show_n=st.get("show_n", False),
                      title=st.get("title", title), xlabel=st.get("xlabel", xlabel),
                      metric=metric, unit_note=note, ymax=st.get("ymax"))
    make_boxplot(conds, data, os.path.join(outdir, "uptake_box.png"),
                 metric=metric, title=st.get("title", title))
    xlsx = write_excel(plot_df, metric, summary, stats, n_comp,
                       os.path.join(outdir, "Results.xlsx"),
                       error=st.get("error", error), methods=meth,
                       simple_tables=simple_tables, summary_full=summary_full,
                       scope=scope,
                       stats_full_df=(pd.DataFrame(stats_full) if stats_full
                                      else None))

    print(f"\n[report] mean +/- {st.get('error', error).upper()}, "
          f"each point = one {st.get('level', level)}")
    for r in summary:
        flag = "   [no dye - not plotted]" if r["condition"] in dropped else ""
        print(f"    {r['condition'][:38]:38s} n={r['n_cells']:4d} cells / "
              f"{r['n_images']:2d} images   mean={r['mean']:>12,.0f} "
              f"+/- {r['sd']:>11,.0f}{flag}")
    if stats and scope == "top" and stats_full:
        by_pair = {frozenset((r["condition_a"], r["condition_b"])): r
                   for r in stats_full}
        clashes = []
        for r in stats:
            o = by_pair.get(frozenset((r["condition_a"], r["condition_b"])))
            if o and (r["stars_cells"] == "ns") != (o["stars_cells"] == "ns"):
                clashes.append((r, o))
        if clashes:
            print("\n[report] !! THE FIGURE AND THE FULL POPULATION DISAGREE !!")
            for shown, allc in clashes:
                print(f"    {shown['condition_a'][:26]} vs {shown['condition_b'][:26]}")
                print(f"        the {int(shown['n_cells_a'])} cells plotted : "
                      f"p={shown['p_cells_bonferroni']:.2e}  {shown['stars_cells']}")
                print(f"        all {int(allc['n_cells_a'])} vs "
                      f"{int(allc['n_cells_b'])} cells   : "
                      f"p={allc['p_cells_bonferroni']:.2e}  {allc['stars_cells']}")
            print("    Keeping only the brightest N compares the tails of each")
            print("    group, which can hide a real difference between the")
            print("    populations. Quote the full-population result, or plot")
            print("    every cell (Edit Figure -> Data -> plot all).")

    if stats:
        print(f"\n[report] pairwise (Bonferroni x{n_comp}):  per cell | per image")
        for s in stats:
            pi = ("%.2e" % s["p_images_bonferroni"]
                  if s["p_images_bonferroni"] == s["p_images_bonferroni"] else "n/a")
            warn = ("   <-- disagrees; quote the image-level value"
                    if s["stars_cells"] != "ns" and s["stars_images"] == "ns" else "")
            print(f"    {s['condition_a'][:22]:22s} vs {s['condition_b'][:22]:22s} "
                  f"{s['p_cells_bonferroni']:.2e} {s['stars_cells']:3s} | "
                  f"{pi} {s['stars_images']:3s}{warn}")
    print(f"[report] methods text -> {os.path.join(outdir, 'methods.txt')}")
    return fig, xlsx


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    kw = dict(metric="cy3_integrated", error="sd", level="cell",
              title="", xlabel="", show_stars=True)
    for a in sys.argv[1:]:
        for k in ("metric", "error", "level", "title", "xlabel"):
            if a.startswith(f"--{k}="):
                kw[k] = a.split("=", 1)[1]
        if a == "--no-stars":
            kw["show_stars"] = False
    if not args:
        print('Usage: python -m cell_viability.report "<results folder>" '
              '[--metric=cy3_integrated|cy3_mean] [--error=sd|sem|ci95] '
              '[--level=cell|image] [--title="HEK 293"] [--no-stars]')
        sys.exit(1)
    build_report(args[0], **kw)

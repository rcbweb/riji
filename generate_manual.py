import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title=None, border_color="2B6CB0", bg_color="F0F4F8"):
    """Adds a stylish callout box to the document."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.bold = True
        run_title.font.name = "Segoe UI"
        run_title.font.size = Pt(10.5)
        run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        
    run_text = p.add_run(text)
    run_text.font.name = "Segoe UI"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    # Spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def format_table_header(row, col_widths, bg_color="1B365D"):
    for idx, cell in enumerate(row.cells):
        cell.width = col_widths[idx]
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.bold = True
                run.font.name = "Segoe UI"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

def format_table_row(row, col_widths, is_even=False):
    bg_color = "F7FAFC" if is_even else "FFFFFF"
    for idx, cell in enumerate(row.cells):
        cell.width = col_widths[idx]
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.name = "Segoe UI"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

def build_manual():
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
    # Styles helper
    def add_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(16)
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        return p

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = "Segoe UI"
            r_pre.font.size = Pt(10)
            r_pre.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = "Segoe UI"
            r_pre.font.size = Pt(10)
            r_pre.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        return p

    # --- Title Section ---
    add_title("RIJI: Researcher User Manual & Step-by-Step Guide")
    add_subtitle("Desktop Colocalization, Live-Cell Viability, and Quantitative Microscopy Analysis\nAuthor: Rchin Bari  |  Software Version 1.0.0  |  Open Source Release")

    add_callout(
        doc,
        "Riji is designed for biological and biophysical researchers who need fast, objective, "
        "and fully reproducible image analysis without writing code. This manual walks you through "
        "folder setup, navigating the user interface, inspecting segmented cells, customizing figures, "
        "and understanding the auto-generated Excel and Methods outputs.",
        title="Quick Reference Summary"
    )

    # --- Section 1: Quick Launch ---
    add_h1("1. How to Launch Riji")
    add_p("Riji is shipped as a self-contained desktop program for both Windows and macOS:")
    
    add_h2("On Windows:")
    add_bullet(" Extract the downloaded Riji_Windows.zip file completely (Right-click > Extract All). Do not run Riji directly from inside the compressed zip folder.", "1. Extract Properly:")
    add_bullet(" Double-click Riji.exe in the extracted folder. No installation, Python, or admin rights are required.", "2. Launch:")
    add_bullet(" If Windows SmartScreen displays 'Windows protected your PC', click 'More info' and then 'Run anyway'. This occurs only on the first launch because the software is unsigned academic freeware (see SECURITY.md).", "3. SmartScreen:")

    add_h2("On macOS:")
    add_bullet(" Double-click Riji-Mac.zip and drag the Riji-Mac folder onto your Desktop.", "1. Unpack:")
    add_bullet(" Open Terminal (Cmd+Space, type Terminal, hit Enter), paste: cd ~/Desktop/Riji-Mac && bash install.command and press Enter.", "2. First-Time Setup:")
    add_bullet(" The installer sets up a private environment inside the folder and downloads the Cellpose segmentation model. Once finished, double-click Riji.app to run any time.", "3. Launch:")

    # --- Section 2: Folder Organization ---
    add_h1("2. How to Organize Your Microscopy Images")
    add_p("Riji requires an organized folder structure. It automatically treats every subfolder as a distinct experimental condition, and every image within that subfolder as a replicate:")

    add_callout(
        doc,
        "My Experiment Folder/             <-- POINT RIJI AT THIS TOP-LEVEL FOLDER\n"
        "   ├── Control/\n"
        "   │     ├── wellA01_01.czi (or .nd2, .tif)\n"
        "   │     └── wellA01_02.czi\n"
        "   ├── Treatment A/\n"
        "   │     ├── wellB01_01.czi\n"
        "   │     └── wellB01_02.czi\n"
        "   └── Treatment B/\n"
        "         ├── wellC01_01.czi\n"
        "         └── wellC01_02.czi",
        title="Recommended Folder Hierarchy",
        border_color="1B365D",
        bg_color="F8FAFC"
    )

    add_bullet(" Supported Formats: Zeiss (.czi), Nikon (.nd2), and multi-page OME-TIFF (.tif, .tiff).", "File Formats:")
    add_bullet(" Condition Subfolders: Name your subfolders clearly (e.g. 'Control', '10uM Drug', '30uM Drug'). These names will automatically label the bars on your final figures and Excel sheets.", "Condition Names:")
    add_bullet(" Subfolder Depth: Keep image files directly inside their condition folder (one level deep).", "Nesting:")

    # --- Section 3: UI Tour ---
    add_h1("3. Step-by-Step UI Walkthrough")
    add_p("When you launch Riji, the main window presents a straightforward, top-to-bottom workflow:")

    add_h2("Step 3.1: Selecting Image and Output Paths")
    add_bullet(" Click 'Choose folder...' and select the top-level experiment folder containing your condition subfolders.", "Choose folder:")
    add_bullet(" By default, Riji creates a folder named coloc_results inside your selected image directory. If you want results saved elsewhere, click 'Save results to...'.", "Save results to:")

    add_h2("Step 3.2: The Action Buttons (The Core Pipeline)")
    add_p("Below the folder pickers is the primary control bar containing the numbered analysis buttons:")

    # Action Buttons Table
    tbl_actions = doc.add_table(rows=6, cols=3)
    tbl_actions.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(1.8), Inches(1.8), Inches(2.9)]
    
    hdr = tbl_actions.rows[0]
    hdr.cells[0].paragraphs[0].text = "Button"
    hdr.cells[1].paragraphs[0].text = "Purpose"
    hdr.cells[2].paragraphs[0].text = "When to Use"
    format_table_header(hdr, widths)
    
    actions_data = [
        ("1. Find and measure cells", "Automated cell segmentation & dye uptake quantification", "Always run this first for live-cell uptake experiments. Uses Cellpose on brightfield images."),
        ("2. Check the cells", "Interactive visual inspection & cell editing tool", "Run after Step 1 if you want to visually verify cell outlines, drop out-of-focus cells, or reject debris."),
        ("3. Adjust the figure", "Publication figure customization editor", "Use to rename treatment bars, reorder conditions, change bar colors, or hide unwanted replicates."),
        ("Colocalization in cells", "Single-cell multi-channel overlap", "Measures Pearson (r) and Manders (M1, M2) overlap strictly inside segmented cell boundaries."),
        ("Colocalization (whole image)", "Whole-frame multi-channel overlap", "Classic whole-frame colocalization across entire FOVs. Requires dye names entered in settings.")
    ]
    
    for idx, (b_name, b_purp, b_when) in enumerate(actions_data):
        row = tbl_actions.rows[idx + 1]
        row.cells[0].paragraphs[0].text = b_name
        row.cells[1].paragraphs[0].text = b_purp
        row.cells[2].paragraphs[0].text = b_when
        format_table_row(row, widths, is_even=(idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- Section 4: Deep Dive Step 1 ---
    add_h1("4. Running '1. Find and Measure Cells'")
    add_p("Clicking '1.  Find and measure cells' opens a parameters dialog:")
    add_bullet(" Set how many cells you want measured from each experimental folder (default is 20). Riji selects the highest-quality, most viable cells first.", "Cells per folder (N):")
    add_bullet(" Choose 'Total intensity (integrated)' to measure total dye payload, or 'Mean intensity' to measure concentration/brightness per unit area.", "Intensity Metric:")
    add_bullet(" Default is 'Live cells'. Select 'Dead cells' only if your study specifically quantifies membrane-permeabilized or dead populations.", "Target Population:")
    add_bullet(" 'Separated cells' (sparse cultures) or 'Confluent sheet' (monolayers with touching cell borders). 'Auto' detects automatically.", "Cell Culture Layout:")
    add_bullet(" Check 'Equalise depth across conditions' if you want every condition to contribute an identical number of cells to statistical figures.", "Depth Equalisation:")
    add_bullet(" Press START. The first run analyzes every image and caches segmentation for rapid future access. When complete, click 'Open Results Folder'.", "Execution:")

    # --- Section 5: Deep Dive Step 2 (Check Cells) ---
    add_h1("5. Inspecting Cells: '2. Check the Cells' (Interactive Editor)")
    add_p("Riji features human-in-the-loop validation. Press '2. Check the cells' to open the visual reviewer. Every identified object is outlined with an intuitive color code:")

    # Colors Table
    tbl_colors = doc.add_table(rows=5, cols=3)
    tbl_colors.alignment = WD_TABLE_ALIGNMENT.CENTER
    w_col = [Inches(1.2), Inches(1.8), Inches(3.5)]
    
    hdr_c = tbl_colors.rows[0]
    hdr_c.cells[0].paragraphs[0].text = "Outline Color"
    hdr_c.cells[1].paragraphs[0].text = "Meaning"
    hdr_c.cells[2].paragraphs[0].text = "Explanation & Action"
    format_table_header(hdr_c, w_col)
    
    colors_data = [
        ("Green", "Measured Cell", "A valid live cell included in your current N selection and statistics."),
        ("Red", "Rejected Cell", "A real cell that was excluded (e.g. ranked below top N, touching edge, or out of focus). Click to include it."),
        ("Blue", "Hand-Drawn Cell", "A cell manually outlined by you using the 'D' key."),
        ("Grey", "Not A Cell (Debris)", "Classified as debris, smear, or dust. Completely removed from cell counts and uptake stats.")
    ]
    for idx, (c_name, c_mean, c_act) in enumerate(colors_data):
        row = tbl_colors.rows[idx + 1]
        row.cells[0].paragraphs[0].text = c_name
        row.cells[1].paragraphs[0].text = c_mean
        row.cells[2].paragraphs[0].text = c_act
        format_table_row(row, w_col, is_even=(idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("Keyboard and Mouse Controls in the Reviewer:")
    
    # Keyboard Table
    tbl_keys = doc.add_table(rows=6, cols=2)
    tbl_keys.alignment = WD_TABLE_ALIGNMENT.CENTER
    w_keys = [Inches(2.0), Inches(4.5)]
    
    hdr_k = tbl_keys.rows[0]
    hdr_k.cells[0].paragraphs[0].text = "Key / Mouse Action"
    hdr_k.cells[1].paragraphs[0].text = "Action Performed"
    format_table_header(hdr_k, w_keys)
    
    keys_data = [
        ("Left-Click on a cell", "Toggle Keep / Drop. Switches between Green (measured) and Red (rejected)."),
        ("Press 'D'", "Draw mode. Click points around an unsegmented cell to outline it manually (turns Blue)."),
        ("Right-Click or 'X'", "Mark as 'Not a cell'. Turns Grey. Strips debris and smears from all counts permanently."),
        ("'N' / 'P'", "Navigate to Next image ('N') or Previous image ('P')."),
        ("Press 'R'", "Rebuild and Retrain. Recalculates all figures and spreadsheets with your edits, and updates the AI selection model.")
    ]
    for idx, (k_name, k_act) in enumerate(keys_data):
        row = tbl_keys.rows[idx + 1]
        row.cells[0].paragraphs[0].text = k_name
        row.cells[1].paragraphs[0].text = k_act
        format_table_row(row, w_keys, is_even=(idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- Section 6: Figure Customizer ---
    add_h1("6. Customizing Publication Figures: '3. Adjust the Figure'")
    add_p("You never need to copy-paste data into Prism or Origin just to format a plot. Pressing '3. Adjust the figure' opens the Figure Customizer:")
    add_bullet(" Click any condition name to edit the label displayed on the x-axis (e.g. change 'treatment1' to '5 uM Compound X').", "Rename Condition Bars:")
    add_bullet(" Drag or use up/down arrows to order your conditions (e.g. Control first, followed by ascending doses).", "Reorder Bars:")
    add_bullet(" Click on a condition's color swatch to pick exact RGB or hex colors matching your journal or lab palette.", "Custom Colors:")
    add_bullet(" If a specific field of view had an air bubble or focal drift, uncheck it here. It is excluded without deleting the raw file.", "Hide Bad Images:")
    add_bullet(" The vector PDF/PNG figures and Results.xlsx workbook are regenerated simultaneously, ensuring figures and tables always match 100%.", "Synchronized Updates:")

    # --- Section 7: Output Files ---
    add_h1("7. Understanding Your Output Files")
    add_p("All output files are saved into the coloc_results/ folder:")
    
    add_h2("1. Figures and Tables/Results.xlsx")
    add_bullet(" Lists every individual measured cell, its parent image, shape features, area, and exact background-subtracted intensity values.", "'Cells' Sheet:")
    add_bullet(" Condition means, medians, standard deviations, and standard errors of the mean (SEM).", "'Summary' Sheet:")
    add_bullet(" Statistical tests reported TWO ways: Per-cell and Per-image. Because cells in the same image are not independent, per-image statistics provide the most rigorous metric for peer review.", "'Statistics' Sheet:")

    add_h2("2. Figures and Tables/methods.txt (Manuscript Ready)")
    add_p("Riji automatically writes the complete microscopy analysis methods paragraph for your manuscript, populated with your exact numbers:")
    add_callout(
        doc,
        "\"Confocal fluorescence micrographs were analyzed using Riji (v1.0.0). Cellular boundaries were segmented from brightfield transmission channels using deep learning (Cellpose) independently of fluorescence intensity to prevent dead-cell uptake bias. Fluorescent dye intensities were quantified across N=20 cells per condition after subtracting local annular background... Replicate variability was evaluated across independent fields of view...\"",
        title="Example Generated Methods Text",
        border_color="2B6CB0",
        bg_color="F0F4F8"
    )

    add_h2("3. Figures and Tables/uptake_figure.pdf and .png")
    add_bullet(" High-resolution vector PDF suitable for Adobe Illustrator / publication submission, plus a 300 DPI PNG preview.", "Publication Graphics:")

    add_h2("4. Cell Images/ (Audit Montages)")
    add_bullet(" Contains image montages for every condition, displaying every cell selected, along with full-field overlay masks showing green/red/blue selections for complete transparency.", "Full Visual Audit:")

    # --- Section 8: Scientific Best Practices ---
    add_h1("8. Key Scientific Design Principles in Riji")
    add_bullet(" In live-cell dye assays, dead or dying cells frequently exhibit compromised membranes that hyper-accumulate dye. Naive thresholding will identify dead cells as the 'highest uptake' hits. Riji identifies cells based on brightfield morphology FIRST, and only measures fluorescence within confirmed viable cells.", "Why Brightness Never Decides Cell Identity:")
    add_bullet(" Cells in the same field of view share laser power, focus, and local cell density. Riji calculates statistics both ways, ensuring you have robust replication numbers for reviewers.", "Per-Cell vs Per-Image Statistics:")
    add_bullet(" Every analysis run saves a snapshot of the parameters, models, and code in the results folder, ensuring any run can be audited years later.", "Reproducibility Guarantee:")

    # --- Save Document ---
    out_path = os.path.join(os.getcwd(), "Riji_User_Manual.docx")
    doc.save(out_path)
    print(f"User manual generated successfully at: {out_path}")

if __name__ == "__main__":
    build_manual()
